import argparse
import hashlib
import json
import re
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import uuid4

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from core.database import SessionLocal, User, init_db
from core.config import FACE_OSS_BUCKET_NAME, FACE_OSS_ENDPOINT
from core.oss_utils import upload_bytes_to_oss


DEFAULT_EXCEL_FILE = PROJECT_ROOT / "信息.xlsx"
INITIAL_PASSWORD_SUFFIX = "2026"
REQUIRED_HEADERS = ["姓名", "护照号码", "注册账号", "人脸照片"]


@dataclass
class ImportStats:
    total: int = 0
    created: int = 0
    updated: int = 0
    skipped: int = 0
    failed: int = 0


def clean_cell(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def hash_password(password: str) -> str:
    return hashlib.sha256(password.encode("utf-8")).hexdigest()


def normalize_email(value) -> str:
    return clean_cell(value).lower()


def normalize_passport_no(value) -> str:
    return re.sub(r"\s+", "", clean_cell(value)).upper()


def password_from_passport(passport_no: str) -> str:
    passport = normalize_passport_no(passport_no)
    if not passport:
        raise ValueError("护照号码为空，无法生成初始密码")
    return f"{passport}{INITIAL_PASSWORD_SUFFIX}"


def load_workbook(path: Path):
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，请先执行：pip install openpyxl") from exc
    return openpyxl.load_workbook(path, read_only=False, data_only=True)


def get_header_map(sheet) -> dict[str, int]:
    headers = {}
    for index, value in enumerate(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)), start=1):
        header = clean_cell(value)
        if header:
            headers[header] = index
    return headers


def image_extension(image) -> str:
    ext = (getattr(image, "format", None) or "jpeg").lower()
    if ext == "jpeg":
        return "jpg"
    if ext in {"jpg", "png", "webp"}:
        return ext
    return "jpg"


def get_images_by_row(sheet, image_column: int) -> dict[int, object]:
    images_by_row = {}
    for image in getattr(sheet, "_images", []):
        anchor = getattr(image, "anchor", None)
        marker = getattr(anchor, "_from", None)
        if not marker:
            continue
        row = marker.row + 1
        col = marker.col + 1
        if col == image_column:
            images_by_row[row] = image
    return images_by_row


def read_candidates_from_excel(path: Path) -> list[dict]:
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = get_header_map(sheet)
    missing = [name for name in REQUIRED_HEADERS if name not in headers]
    if missing:
        raise ValueError(f"Excel 缺少必需列：{', '.join(missing)}")

    images_by_row = get_images_by_row(sheet, headers["人脸照片"])
    rows = []
    for row_index in range(2, sheet.max_row + 1):
        real_name = clean_cell(sheet.cell(row_index, headers["姓名"]).value)
        passport_no = normalize_passport_no(sheet.cell(row_index, headers["护照号码"]).value)
        email = normalize_email(sheet.cell(row_index, headers["注册账号"]).value)
        remark = clean_cell(sheet.cell(row_index, headers["备注"]).value) if "备注" in headers else ""

        if not real_name and not passport_no and not email:
            continue

        rows.append(
            {
                "row": row_index,
                "real_name": real_name,
                "passport_no": passport_no,
                "email": email,
                "password": password_from_passport(passport_no) if passport_no else "",
                "remark": remark,
                "image": images_by_row.get(row_index),
            }
        )
    return rows


def validate_candidates(rows: list[dict]) -> list[dict]:
    errors = []
    seen_emails = {}
    seen_passports = {}

    for row in rows:
        row_errors = []
        email = row["email"]
        passport_no = row["passport_no"]

        if not row["real_name"]:
            row_errors.append("姓名为空")
        if not passport_no:
            row_errors.append("护照号码为空")
        if not email:
            row_errors.append("注册账号/邮箱为空")
        if not row["image"]:
            row_errors.append("人脸照片为空或未内嵌在对应行")

        if email:
            if email in seen_emails:
                row_errors.append(f"邮箱重复，首次出现在第 {seen_emails[email]} 行")
            seen_emails[email] = row["row"]
        if passport_no:
            if passport_no in seen_passports:
                row_errors.append(f"护照号码重复，首次出现在第 {seen_passports[passport_no]} 行")
            seen_passports[passport_no] = row["row"]

        if row_errors:
            errors.append({"row": row["row"], "email": email, "errors": row_errors})

    return errors


def upload_candidate_face(candidate: dict) -> tuple[str, str]:
    image = candidate["image"]
    ext = image_extension(image)
    oss_key = f"face/candidates/{candidate['passport_no']}/{uuid4().hex}.{ext}"
    upload_bytes_to_oss(
        image._data(),
        oss_key,
        True,
        endpoint=FACE_OSS_ENDPOINT,
        bucket_name=FACE_OSS_BUCKET_NAME,
    )
    return oss_key, ext


def upsert_candidate_after_upload(db, candidate: dict, oss_key: str, source: str) -> dict:
    now = datetime.now()
    email = candidate["email"]
    passport_no = candidate["passport_no"]
    real_name = candidate["real_name"]
    password = candidate["password"]

    user = db.query(User).filter(User.passport_no == passport_no).first()
    created = user is None
    if created:
        user = User(
            email=email,
            passport_no=passport_no,
            real_name=real_name,
            password_hash=hash_password(password),
            status="active",
            remark=candidate["remark"],
            source=source,
            imported_at=now,
        )
        db.add(user)
        db.flush()
    else:
        user.email = email
        user.passport_no = passport_no
        user.real_name = real_name
        user.password_hash = hash_password(password)
        user.status = "active"
        user.remark = candidate["remark"]
        user.source = source
        user.imported_at = now

    # Login uses passport number/password. Face verification uses the stored OSS key below.
    user.face_image_oss_key = oss_key
    user.face_enrolled_at = now

    db.commit()
    db.refresh(user)
    return {
        "action": "created" if created else "updated",
        "id": user.id,
        "email": user.email,
        "passport_no": user.passport_no or "",
        "password": password,
        "face_image_oss_key": user.face_image_oss_key or "",
        "status": user.status,
    }


def import_candidates(args) -> dict:
    input_path = Path(args.file).expanduser().resolve()
    if not input_path.is_file():
        raise FileNotFoundError(str(input_path))

    rows = read_candidates_from_excel(input_path)
    row_errors = validate_candidates(rows)
    stats = ImportStats(total=len(rows))
    report = {
        "file": str(input_path),
        "dry_run": args.dry_run,
        "started_at": datetime.now().isoformat(timespec="seconds"),
        "stats": stats.__dict__,
        "row_errors": row_errors,
        "results": [],
    }

    if row_errors:
        stats.failed = len(row_errors)
        report["stats"] = stats.__dict__
        return report

    if args.dry_run:
        stats.skipped = len(rows)
        report["stats"] = stats.__dict__
        report["results"] = [
            {
                "row": row["row"],
                "email": row["email"],
                "passport_no": row["passport_no"],
                "password": row["password"],
                "has_face_image": bool(row["image"]),
            }
            for row in rows
        ]
        return report

    init_db()
    db = SessionLocal()
    try:
        for candidate in rows:
            try:
                oss_key, _ = upload_candidate_face(candidate)
                result = upsert_candidate_after_upload(db, candidate, oss_key, args.source)
                if result["action"] == "created":
                    stats.created += 1
                else:
                    stats.updated += 1
                report["results"].append({"row": candidate["row"], **result})
            except Exception as exc:
                db.rollback()
                stats.failed += 1
                report["row_errors"].append(
                    {"row": candidate["row"], "email": candidate["email"], "errors": [str(exc)]}
                )
                if args.stop_on_error:
                    break
        report["stats"] = stats.__dict__
        return report
    finally:
        db.close()


def write_json_report(path: Path, report: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def print_summary(report: dict) -> None:
    stats = report["stats"]
    print("Candidate import summary:")
    print(f"  file: {report['file']}")
    print(f"  dry_run: {report['dry_run']}")
    print(f"  total: {stats['total']}")
    print(f"  created: {stats['created']}")
    print(f"  updated: {stats['updated']}")
    print(f"  skipped: {stats['skipped']}")
    print(f"  failed: {stats['failed']}")
    if report["row_errors"]:
        print("Row errors:")
        for item in report["row_errors"][:20]:
            print(f"  row {item['row']} email={item.get('email', '')}: {'; '.join(item['errors'])}")
        if len(report["row_errors"]) > 20:
            print(f"  ... {len(report['row_errors']) - 20} more errors")


def parse_args():
    parser = argparse.ArgumentParser(
        description="Import candidates from 信息.xlsx. Login account is passport number; password is uppercase passport number plus 2026."
    )
    parser.add_argument("--file", default=str(DEFAULT_EXCEL_FILE), help="Excel file path. Default: 项目根目录/信息.xlsx")
    parser.add_argument("--dry-run", action="store_true", help="Validate only. Do not upload OSS or write DB.")
    parser.add_argument("--source", default="info_xlsx_import", help="Value written to users.source.")
    parser.add_argument("--stop-on-error", action="store_true", help="Stop after first import error.")
    parser.add_argument("--report", default="", help="Optional JSON report path.")
    return parser.parse_args()


if __name__ == "__main__":
    parsed_args = parse_args()
    import_report = import_candidates(parsed_args)
    print_summary(import_report)
    if parsed_args.report:
        write_json_report(Path(parsed_args.report).expanduser().resolve(), import_report)
