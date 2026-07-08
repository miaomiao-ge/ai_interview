import argparse
import hashlib
import json
import sys
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

DEFAULT_EXCEL_FILE = PROJECT_ROOT / "questions.xlsx"
REQUIRED_HEADERS = ["问题类型", "序号", "题目（中文）", "题目（English）"]
QuestionBank = None
SessionLocal = None
init_db = None


def load_database_api() -> None:
    global QuestionBank, SessionLocal, init_db
    from core import database

    QuestionBank = QuestionBank or database.QuestionBank
    SessionLocal = SessionLocal or database.SessionLocal
    init_db = init_db or database.init_db


@dataclass
class QuestionRow:
    category: str
    question_no: int | None
    content: str
    content_en: str
    source_file: str
    source_row: int
    source_hash: str


@dataclass
class ImportStats:
    total: int = 0
    created: int = 0
    updated: int = 0
    skipped: int = 0
    failed: int = 0


def clean_cell(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def load_workbook(path: Path):
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，请先执行：pip install openpyxl") from exc
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def get_header_map(sheet) -> dict[str, int]:
    rows = sheet.iter_rows(min_row=1, max_row=1, values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return {}
    return {clean_cell(value): index for index, value in enumerate(header_row, start=1) if clean_cell(value)}


def parse_question_no(value) -> int | None:
    cleaned = clean_cell(value)
    if not cleaned:
        return None
    try:
        return int(float(cleaned))
    except ValueError as exc:
        raise ValueError(f"题目序号不是数字：{cleaned}") from exc


def build_source_hash(category: str, question_no: int | None, content: str, content_en: str) -> str:
    raw = "\0".join([category, str(question_no or ""), content, content_en])
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def read_questions_from_excel(path: Path) -> list[QuestionRow]:
    if not path.exists():
        raise FileNotFoundError(f"题库文件不存在：{path}")

    workbook = load_workbook(path)
    sheet = workbook.active
    headers = get_header_map(sheet)
    missing = [name for name in REQUIRED_HEADERS if name not in headers]
    if missing:
        raise ValueError(f"Excel 缺少必需列：{', '.join(missing)}")

    rows: list[QuestionRow] = []
    current_category = ""
    for row_index in range(2, sheet.max_row + 1):
        category_cell = clean_cell(sheet.cell(row_index, headers["问题类型"]).value)
        if category_cell:
            current_category = category_cell

        question_no = parse_question_no(sheet.cell(row_index, headers["序号"]).value)
        content = clean_cell(sheet.cell(row_index, headers["题目（中文）"]).value)
        content_en = clean_cell(sheet.cell(row_index, headers["题目（English）"]).value)

        if not current_category and not content and not content_en:
            continue
        if not current_category:
            raise ValueError(f"第 {row_index} 行缺少问题类型")
        if not content:
            raise ValueError(f"第 {row_index} 行缺少中文题目")

        rows.append(
            QuestionRow(
                category=current_category,
                question_no=question_no,
                content=content,
                content_en=content_en,
                source_file=path.name,
                source_row=row_index,
                source_hash=build_source_hash(current_category, question_no, content, content_en),
            )
        )

    return rows


def find_existing_question(db, row: QuestionRow):
    if QuestionBank is None:
        load_database_api()
    query = db.query(QuestionBank)
    if row.question_no is not None:
        existing = (
            query.filter(
                QuestionBank.category == row.category,
                QuestionBank.question_no == row.question_no,
            )
            .order_by(QuestionBank.id.asc())
            .first()
        )
        if existing:
            return existing
    return (
        query.filter(
            QuestionBank.category == row.category,
            QuestionBank.content == row.content,
        )
        .order_by(QuestionBank.id.asc())
        .first()
    )


def upsert_questions(rows: list[QuestionRow], dry_run: bool = False) -> ImportStats:
    if SessionLocal is None:
        load_database_api()
    stats = ImportStats(total=len(rows))
    db = SessionLocal()
    try:
        now = datetime.now()
        for row in rows:
            existing = find_existing_question(db, row)
            if existing:
                if (
                    getattr(existing, "question_no", None) == row.question_no
                    and (existing.content or "") == row.content
                    and (getattr(existing, "content_en", "") or "") == row.content_en
                    and (getattr(existing, "source_hash", "") or "") == row.source_hash
                ):
                    stats.skipped += 1
                    continue

                existing.category = row.category
                existing.question_no = row.question_no
                existing.content = row.content
                existing.content_en = row.content_en
                existing.source_file = row.source_file
                existing.source_row = row.source_row
                existing.source_hash = row.source_hash
                existing.updated_at = now
                stats.updated += 1
                continue

            db.add(
                QuestionBank(
                    category=row.category,
                    question_no=row.question_no,
                    content=row.content,
                    content_en=row.content_en,
                    source_file=row.source_file,
                    source_row=row.source_row,
                    source_hash=row.source_hash,
                    updated_at=now,
                )
            )
            stats.created += 1

        if dry_run:
            db.rollback()
        else:
            db.commit()
        return stats
    except Exception:
        db.rollback()
        stats.failed += 1
        raise
    finally:
        db.close()


def parse_args():
    parser = argparse.ArgumentParser(description="导入 questions.xlsx 到 question_bank 表")
    parser.add_argument("--file", default=str(DEFAULT_EXCEL_FILE), help="题库 Excel 文件路径")
    parser.add_argument("--dry-run", action="store_true", help="解析并模拟导入，不提交数据库")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    path = Path(args.file).resolve()
    load_database_api()
    init_db()
    rows = read_questions_from_excel(path)
    stats = upsert_questions(rows, dry_run=args.dry_run)
    print(
        json.dumps(
            {
                "status": "completed",
                "file": str(path),
                "dry_run": args.dry_run,
                "stats": asdict(stats),
            },
            ensure_ascii=False,
            indent=2,
        ),
        flush=True,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
